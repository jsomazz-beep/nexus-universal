from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "docs" / "Guia_Uso_WebApp_Strategic_Intelligence.pptx"


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.add_paragraph() if idx else tf.paragraphs[0]
        p.text = bullet
        p.level = 0
        p.font.size = Pt(22)


def add_two_column_slide(prs: Presentation, title: str, left_title: str, left_items: list[str], right_title: str, right_items: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    left = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(6.0), Inches(4.8))
    right = slide.shapes.add_textbox(Inches(6.8), Inches(1.4), Inches(6.0), Inches(4.8))

    ltf = left.text_frame
    ltf.clear()
    p = ltf.paragraphs[0]
    p.text = left_title
    p.font.bold = True
    p.font.size = Pt(24)
    for item in left_items:
        q = ltf.add_paragraph()
        q.text = f"- {item}"
        q.font.size = Pt(20)

    rtf = right.text_frame
    rtf.clear()
    p2 = rtf.paragraphs[0]
    p2.text = right_title
    p2.font.bold = True
    p2.font.size = Pt(24)
    for item in right_items:
        q2 = rtf.add_paragraph()
        q2.text = f"- {item}"
        q2.font.size = Pt(20)


def style_slides(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Segoe UI"
                    if paragraph.level == 0 and shape == slide.shapes.title:
                        run.font.color.rgb = RGBColor(15, 23, 42)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Guia de Uso do WebApp",
        "Strategic Intelligence Reporter | Versão prática para operação diária",
    )

    add_bullets_slide(
        prs,
        "1. Objetivo do WebApp",
        [
            "Gerar relatórios de notícias e oportunidades de forma rápida e padronizada.",
            "Permitir configuração por categorias, países, fontes e período.",
            "Salvar perfis de busca para execuções recorrentes sem retrabalho.",
        ],
    )

    add_bullets_slide(
        prs,
        "2. Como abrir o sistema",
        [
            "Executar: dist\\StrategicIntelligenceWebApp.exe",
            "O navegador abre em: http://127.0.0.1:8787",
            "Se bloqueado pelo Windows: Mais informações > Executar assim mesmo.",
        ],
    )

    add_bullets_slide(
        prs,
        "3. Passo a passo rápido",
        [
            "Selecionar categorias/palavras-chave.",
            "Selecionar países/regiões + subregiões opcionais.",
            "Configurar quantidade de notícias e intervalo em dias.",
            "Ajustar FONTES e clicar em Gerar Relatório.",
        ],
    )

    add_two_column_slide(
        prs,
        "4. Seção FONTES (organizada)",
        "Blocos recolhíveis",
        [
            "Fontes configuradas para coleta",
            "Fontes extras sugeridas (one-click)",
            "Pacotes geográficos",
            "Novas fontes por URL (RSS/feed)",
        ],
        "Como usar melhor",
        [
            "Marque apenas fontes aderentes ao tema",
            "Use pacotes por país quando precisar cobertura ampla",
            "Adicione RSS próprios no campo de URL",
        ],
    )

    add_bullets_slide(
        prs,
        "5. Pacotes geográficos por país",
        [
            "Cada país possui checkbox individual.",
            "Selecione somente os países necessários para reduzir ruído.",
            "As fontes dos países marcados entram automaticamente na coleta.",
        ],
    )

    add_bullets_slide(
        prs,
        "6. Perfis de busca",
        [
            "Salvar perfil: guarda toda a configuração atual.",
            "Duplicar perfil: cria variações rápidas.",
            "Excluir perfil: remove configurações antigas.",
            "Definir padrão: carrega automaticamente ao abrir.",
        ],
    )

    add_bullets_slide(
        prs,
        "7. Resultado e visualização",
        [
            "Após executar, o relatório abre no iframe da própria tela.",
            "Também é possível abrir em nova aba e escolher relatórios antigos.",
            "Use o status para verificar coletadas, filtradas e aprovadas.",
        ],
    )

    add_bullets_slide(
        prs,
        "8. Dicas para evitar relatório vazio",
        [
            "Aumentar intervalo de dias (ex.: 14 a 30).",
            "Selecionar mais países/fontes relevantes.",
            "Evitar filtros excessivamente restritivos ao mesmo tempo.",
            "Conferir logs quando houver baixo volume.",
        ],
    )

    add_bullets_slide(
        prs,
        "9. Suporte rápido (troubleshooting)",
        [
            "EXE não abre: verificar antivírus/SmartScreen.",
            "Erro no build: fechar EXE anterior e pausar OneDrive na pasta.",
            "Poucos resultados: revisar fontes e categorias multilíngues.",
        ],
    )

    add_bullets_slide(
        prs,
        "10. Fluxo operacional recomendado",
        [
            "1) Carregar perfil padrão.",
            "2) Ajustar países e fontes para o ciclo atual.",
            "3) Gerar relatório.",
            "4) Validar destaques e compartilhar.",
        ],
    )

    style_slides(prs)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(path)
